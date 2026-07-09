import os, sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

OUTPUT_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(OUTPUT_DIR)

OUTPUTS = os.path.join(ROOT, "outputs")
SLIDES_OUT = os.path.join(OUTPUT_DIR, "slides.pptx")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BG = RGBColor(0xF5, 0xF0, 0xE1)
ACCENT = RGBColor(0x2C, 0x3E, 0x50)
TEXT = RGBColor(0x2C, 0x3E, 0x50)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x27, 0xAE, 0x60)
RED = RGBColor(0xE7, 0x4C, 0x3C)
BLUE = RGBColor(0x29, 0x80, 0xB9)

def new_slide():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = BG
    return slide

def title_bar(slide, text, color=ACCENT):
    box = slide.shapes.add_shape(1, Inches(0), Inches(0), prs.slide_width, Inches(1.0))
    box.fill.solid()
    box.fill.fore_color.rgb = color
    box.line.fill.background()
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.5)
    tf.margin_top = Inches(0.15)
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = WHITE

def bullet_list(slide, items, left=0.7, top=1.4, width=11.5, height=5.5, size=20):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(size)
        p.font.color.rgb = TEXT
        p.space_after = Pt(8)
        p.level = 0

def add_image(slide, path, left, top, width, height):
    if os.path.exists(path):
        slide.shapes.add_picture(path, Inches(left), Inches(top), Inches(width), Inches(height))

def table(slide, headers, rows, left, top, width, height, col_size=None):
    tbl = slide.shapes.add_table(len(rows)+1, len(headers), Inches(left), Inches(top), Inches(width), Inches(height)).table
    for i, h in enumerate(headers):
        cell = tbl.cell(0, i)
        cell.text = h
        for par in cell.text_frame.paragraphs:
            par.font.bold = True
            par.font.size = Pt(18)
            par.font.color.rgb = WHITE
            par.alignment = PP_ALIGN.CENTER
        cell.fill.solid()
        cell.fill.fore_color.rgb = ACCENT
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = tbl.cell(r_idx+1, c_idx)
            cell.text = str(val)
            for par in cell.text_frame.paragraphs:
                par.font.size = Pt(16)
                par.font.color.rgb = TEXT
                par.alignment = PP_ALIGN.CENTER
    if col_size:
        for i, s in enumerate(col_size):
            tbl.columns[i].width = Inches(s)

# ====== SLIDE 1: Title ======
sl = new_slide()
box = sl.shapes.add_shape(1, Inches(1), Inches(1.5), Inches(11), Inches(4.5))
box.fill.solid(); box.fill.fore_color.rgb = ACCENT; box.line.fill.background()
tf = box.text_frame; tf.word_wrap = True; tf.margin_left=Inches(0.3); tf.margin_top=Inches(0.2)
p = tf.paragraphs[0]; p.text = "Ứng dụng mạng nơ-ron nhân tạo\nlàm hàm heuristic cho thuật toán A*\ntrong bài toán tìm đường mê cung"
p.font.size = Pt(34); p.font.bold = True; p.font.color.rgb = WHITE; p.alignment = PP_ALIGN.CENTER
p2 = tf.add_paragraph(); p2.text = ""; p2.space_after = Pt(14)
p3 = tf.add_paragraph(); p3.text = "Nhóm Tờ Linh: Nguyễn Quốc Minh – Võ Nguyễn Tấn Tài"
p3.font.size = Pt(18); p3.font.color.rgb = RGBColor(0xD5,0xD8,0xDC); p3.alignment = PP_ALIGN.CENTER
p4 = tf.add_paragraph(); p4.text = "Học kỳ II năm học 2025–2026"
p4.font.size = Pt(16); p4.font.color.rgb = RGBColor(0xD5,0xD8,0xDC); p4.alignment = PP_ALIGN.CENTER

# ====== SLIDE 2: Mục tiêu đề tài ======
sl = new_slide(); title_bar(sl, "Mục tiêu đề tài")
bullet_list(sl, [
    "Hiểu ANN học như thế nào qua forward, loss, backpropagation, epoch.",
    "Phân biệt vai trò của training, validation và testing set.",
    "Phân biệt các chỉ số đánh giá: Accuracy (phân loại) vs. MSE/MAE (hồi quy).",
    "Minh họa underfitting, overfitting và good fit bằng đồ thị train thật.",
    "Tích hợp ANN vào A* để thay thế heuristic Manhattan trong demo mê cung.",
])

# ====== SLIDE 3: Bài toán mê cung ======
sl = new_slide(); title_bar(sl, "Bài toán mê cung")
bullet_list(sl, [
    "Mê cung kích thước 20×20: 0 là đường đi, 1 là tường.",
    "Tác nhân đi bốn hướng: lên, xuống, trái, phải.",
    "Yêu cầu: tìm đường từ Start đến Goal với số bước ngắn nhất.",
])

# ====== SLIDE 4: BFS ======
sl = new_slide(); title_bar(sl, "BFS (Breadth-First Search)")
bullet_list(sl, [
    "Duyệt đồ thị theo từng lớp khoảng cách từ điểm bắt đầu.",
    "Đảm bảo tìm đường ngắn nhất trong mê cung không trọng số.",
    "Nhược điểm: có thể duyệt nhiều node vì không dùng thông tin hướng đến goal.",
    "Trong đề tài, BFS có hai vai trò:",
    "    1) Tạo nhãn true_distance cho dataset (chạy từ goal).",
    "    2) Là thuật toán mốc so sánh trong thực nghiệm tìm đường.",
    "BFS KHÔNG được dùng để sinh map. Map được sinh ngẫu nhiên.",
])

# ====== SLIDE 5: A* và heuristic ======
sl = new_slide(); title_bar(sl, "Thuật toán A* và Heuristic")
bullet_list(sl, [
    "Công thức đánh giá: f(n) = g(n) + h(n)",
    "  – g(n): chi phí đã đi từ start đến node n.",
    "  – h(n): chi phí ước lượng từ n đến goal (heuristic).",
    "  – f(n): tổng chi phí dự đoán.",
    "Heuristic càng gần khoảng cách thật, A* thường duyệt càng ít node.",
])

# ====== SLIDE 6: Manhattan ======
sl = new_slide(); title_bar(sl, "Heuristic Manhattan")
bullet_list(sl, [
    "Công thức: h(n) = |x − goal_x| + |y − goal_y|",
    "Đơn giản, nhanh, phù hợp lưới 4 hướng.",
    "Hạn chế: không học từ dữ liệu, không nhìn thấy vật cản.",
])

# ====== SLIDE 7: Ý tưởng ANN heuristic ======
sl = new_slide(); title_bar(sl, "Ý tưởng: ANN làm heuristic")
bullet_list(sl, [
    "Dùng ANN để dự đoán khoảng cách còn lại đến goal.",
    "ANN nhận: tọa độ, khoảng cách tương đối, tường lân cận (10 đặc trưng).",
    "Giá trị dự đoán được dùng làm h(n) trong A*.",
    "→ Giao thoa giữa học máy (Keras / tf.keras) và thuật toán tìm kiếm.",
])

# ====== SLIDE 8: ANN học thế nào ======
sl = new_slide(); title_bar(sl, "ANN học như thế nào")
bullet_list(sl, [
    "Forward propagation: dữ liệu đi qua mạng để tạo dự đoán.",
    "Loss: đo sai số giữa dự đoán và nhãn thật (MSE).",
    "Backpropagation: truyền sai số ngược để tính gradient.",
    "Optimizer (Adam) cập nhật trọng số nhằm giảm loss.",
    "Epoch: một lần mô hình học qua toàn bộ tập train.",
])

# ====== SLIDE 9: Train / Validation / Test ======
sl = new_slide(); title_bar(sl, "Training, Validation, Testing")
bullet_list(sl, [
    "Training set: dùng để cập nhật trọng số mô hình.",
    "Validation set: theo dõi quá trình học, chọn cấu hình.",
    "Testing set: đánh giá cuối cùng sau khi chọn mô hình.",
    "Đề tài chia 70% – 15% – 15% (70.000 / 15.000 / 15.000 dòng).",
])

# ====== SLIDE 10: Chỉ số đánh giá ======
sl = new_slide(); title_bar(sl, "Chỉ số đánh giá")
bullet_list(sl, [
    "Accuracy và ROC/AUC: dùng cho bài toán phân loại (classification).",
    "Đề tài này là hồi quy (regression) vì ANN dự đoán số bước.",
    "Loss: Mean Squared Error (MSE) – Metric: Mean Absolute Error (MAE).",
    "MAE = 2.284 nghĩa là mô hình sai trung bình ~2.28 bước.",
    "Cross-validation: kiểm tra độ ổn định qua 5 lần chia dữ liệu.",
])

# ====== SLIDE 11: Input / Output ======
sl = new_slide(); title_bar(sl, "Input và Output của ANN")
bullet_list(sl, [
    "Input (10 đặc trưng):",
    "  x, y, goal_x, goal_y, dx, dy, up_wall, down_wall, left_wall, right_wall",
    "Output: true_distance (số bước ngắn nhất đến goal).",
    "Nhãn true_distance do BFS tạo từ mỗi map ngẫu nhiên.",
    "Mô hình hồi quy: lớp output không dùng softmax.",
])

# ====== SLIDE 12: Underfitting ======
sl = new_slide(); title_bar(sl, "Underfitting", RED)
add_image(sl, os.path.join(OUTPUTS, "underfit_loss.png"), 0.3, 1.2, 6.5, 3.8)
bullet_list(sl, [
    "Mạng quá nhỏ: Dense(4) + output, 5 epoch.",
    "Train/Val loss còn cao, MAE chưa tốt.",
    "Test MSE = 16.322, MAE = 2.669.",
], left=7.0, top=1.2, width=6.0, height=5.5, size=18)

# ====== SLIDE 13: Overfitting ======
sl = new_slide(); title_bar(sl, "Overfitting", RED)
bullet_list(sl, [
    "Mạng lớn: 512–512–256–128 neuron, 150 epoch.",
    "Chỉ train trên 3.000 dòng.",
    "Train loss giảm nhưng validation loss tăng.",
    "Test MSE = 24.811, MAE = 3.064.",
], left=7.0, top=1.2, width=6.0, height=5.0, size=18)
add_image(sl, os.path.join(OUTPUTS, "overfit_loss.png"), 0.3, 1.2, 6.5, 3.8)

# ====== SLIDE 14: Good fit ======
sl = new_slide(); title_bar(sl, "Good fit (mô hình chính)", GREEN)
bullet_list(sl, [
    "Kiến trúc: Dense(64) – Dropout(0.2) – Dense(32) – Dense(16).",
    "EarlyStopping dừng ở 20 epoch.",
    "Test MSE = 14.300, MAE = 2.284 (tốt nhất).",
    "Được tích hợp vào A* làm heuristic.",
], left=7.0, top=1.2, width=6.0, height=5.0, size=18)
add_image(sl, os.path.join(OUTPUTS, "goodfit_loss.png"), 0.3, 1.2, 6.5, 3.8)

# ====== SLIDE 15: Cross-validation ======
sl = new_slide(); title_bar(sl, "Cross-validation")
table(sl, ["Fold", "MAE", "MSE"], [
    ["1", "2.250", "15.436"],
    ["2", "2.175", "15.181"],
    ["3", "2.305", "14.340"],
    ["4", "2.206", "14.107"],
    ["5", "2.179", "15.230"],
], left=1.5, top=4.0, width=10.0, height=3.0, col_size=[2, 4, 4])
bullet_list(sl, [
    "5-fold, 10 epoch/fold, 20.000 dòng validation mỗi fold.",
    "Trung bình: MAE = 2.223, MSE = 14.859.",
    "Kết quả cho thấy mô hình good fit ổn định.",
])

# ====== SLIDE 16: Demo so sánh ======
sl = new_slide(); title_bar(sl, "Demo so sánh thuật toán", BLUE)
add_image(sl, os.path.join(OUTPUTS, "demo_path.png"), 0.3, 1.2, 6.5, 4.5)
table(sl, ["Thuật toán", "Đường dài", "Node duyệt", "Thời gian"], [
    ["BFS", "20", "104", "0.09 ms"],
    ["A* + Manhattan", "20", "46", "0.07 ms"],
    ["A* + ANN", "20", "33", "0.48 ms"],
], left=7.0, top=4.0, width=6.0, height=2.5, col_size=[2.0, 1.3, 1.5, 1.2])
bullet_list(sl, [
    "UI tương tác: python -m src.ui",
    "Animate node + đánh số, 3 màu thuật toán.",
], left=7.0, top=2.0, width=6.0, height=2.0, size=16)

# ====== SLIDE 17: Nhận xét ======
sl = new_slide(); title_bar(sl, "Nhận xét và hạn chế")
bullet_list(sl, [
    "A* + ANN duyệt 33 node, ít nhất trong demo (so với 104 của BFS).",
    "ANN chỉ dùng 4 tín hiệu tường lân cận → chưa nhìn toàn bộ map.",
    "Thời gian suy luận ANN (~0.5 ms) > Manhattan (~0.07 ms).",
    "Dataset chỉ có map ngẫu nhiên, chưa bao phủ mọi dạng khó.",
])

# ====== SLIDE 18: Kết luận ======
sl = new_slide(); title_bar(sl, "Kết luận và minh bạch AI")
bullet_list(sl, [
    "Map sinh ngẫu nhiên; BFS dùng để gán nhãn true_distance (không sinh map).",
    "ANN được xây dựng bằng Keras (tf.keras) và tích hợp vào A*.",
    "Ba hiện tượng học minh họa rõ: underfit, overfit, good fit.",
    "Good fit được dùng làm heuristic, giảm node duyệt trong demo.",
    "AI hỗ trợ khung code, debug, tổ chức nội dung.",
    "Nhóm tự chạy lại kết quả, chỉ n cấu hình và giải thích số liệu.",
])

prs.save(SLIDES_OUT)
print(f"Saved: {SLIDES_OUT}")